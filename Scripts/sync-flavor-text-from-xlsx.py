#!/usr/bin/env python3
"""One-time sync of FlavorText.json from the flavor sheet in DEMON_FIGHTER__DATA.xlsx.

Dependency:
    pip install openpyxl

Usage:
    # Dry-run (default) — summary only, no writes
    python Scripts/sync-flavor-text-from-xlsx.py --xlsx /path/to/DEMON_FIGHTER__DATA.xlsx

    # Apply after reviewing dry-run
    python Scripts/sync-flavor-text-from-xlsx.py --xlsx /path/to/DEMON_FIGHTER__DATA.xlsx --write

Patches only:
  - environments.locationDescriptions.<key>
  - environments.roomContexts.<biome>.<roomType>  (sheet key: Forest/boss)
  - combatNarratives.<bank>  (bank column only; key column ignored)

Does not touch the existing PULL_SHEETS pipeline or any other JSON files.
"""

from __future__ import annotations

import argparse
import json
import shutil
import sys
from collections import defaultdict
from pathlib import Path
from typing import Any

try:
    from openpyxl import load_workbook
except ImportError:
    print("Error: openpyxl is required. Install with: pip install openpyxl", file=sys.stderr)
    sys.exit(1)

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_JSON = ROOT / "GameData" / "FlavorText.json"

EXPECTED_LOCATION_DESC_KEYS = 13
EXPECTED_ROOM_CONTEXT_BIOMES = 5
EXPECTED_COMBAT_NARRATIVE_BANKS = 13

SAMPLE_TRUNCATE = 72


def cell_str(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def truncate(text: str, limit: int = SAMPLE_TRUNCATE) -> str:
    text = " ".join(text.split())
    if len(text) <= limit:
        return text
    return text[: limit - 3] + "..."


def find_sheet(workbook, sheet_name: str):
    target = sheet_name.strip().lower()
    for name in workbook.sheetnames:
        if name.strip().lower() == target:
            return workbook[name]
    raise ValueError(
        f"Sheet {sheet_name!r} not found. Available sheets: {', '.join(workbook.sheetnames)}"
    )


def read_header_map(sheet) -> dict[str, int]:
    header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
    mapping: dict[str, int] = {}
    for idx, cell in enumerate(header_row):
        name = cell_str(cell).lower()
        if name:
            mapping[name] = idx
    required = ("section", "bank", "key", "text")
    missing = [col for col in required if col not in mapping]
    if missing:
        raise ValueError(f"Missing required columns: {', '.join(missing)}")
    return mapping


def parse_flavor_sheet(sheet) -> tuple[dict[str, list[str]], dict[str, dict[str, list[str]]], dict[str, list[str]], list[str], int]:
    header = read_header_map(sheet)
    section_idx = header["section"]
    bank_idx = header["bank"]
    key_idx = header["key"]
    text_idx = header["text"]

    location_descriptions: dict[str, list[str]] = defaultdict(list)
    room_contexts: dict[str, dict[str, list[str]]] = defaultdict(lambda: defaultdict(list))
    combat_narratives: dict[str, list[str]] = defaultdict(list)
    errors: list[str] = []
    row_count = 0

    for row_num, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        text = cell_str(row[text_idx] if text_idx < len(row) else None)
        if not text:
            continue

        section = cell_str(row[section_idx] if section_idx < len(row) else None)
        bank = cell_str(row[bank_idx] if bank_idx < len(row) else None)
        key = cell_str(row[key_idx] if key_idx < len(row) else None)
        row_count += 1

        if section == "environments" and bank == "locationDescriptions":
            if not key:
                errors.append(f"Row {row_num}: locationDescriptions row missing key")
                continue
            location_descriptions[key].append(text)
            continue

        if section == "environments" and bank == "roomContexts":
            if "/" not in key:
                errors.append(f"Row {row_num}: roomContexts key must be biome/roomType, got {key!r}")
                continue
            biome, room_type = key.split("/", 1)
            biome = biome.strip()
            room_type = room_type.strip().lower()
            if not biome or not room_type:
                errors.append(f"Row {row_num}: invalid roomContexts key {key!r}")
                continue
            room_contexts[biome][room_type].append(text)
            continue

        if section == "combatNarratives":
            if not bank:
                errors.append(f"Row {row_num}: combatNarratives row missing bank")
                continue
            combat_narratives[bank].append(text)
            continue

        errors.append(
            f"Row {row_num}: unsupported section/bank combination ({section!r}, {bank!r}) — skipped"
        )

    return (
        dict(location_descriptions),
        {biome: dict(room_types) for biome, room_types in room_contexts.items()},
        dict(combat_narratives),
        errors,
        row_count,
    )


def load_json(path: Path) -> dict[str, Any]:
    with path.open(encoding="utf-8") as f:
        return json.load(f)


def ensure_nested(data: dict[str, Any]) -> None:
    environments = data.setdefault("environments", {})
    environments.setdefault("locationDescriptions", {})
    environments.setdefault("roomContexts", {})
    data.setdefault("combatNarratives", {})


def current_location_desc(data: dict[str, Any]) -> dict[str, list[str]]:
    return data.get("environments", {}).get("locationDescriptions", {}) or {}


def current_room_contexts(data: dict[str, Any]) -> dict[str, dict[str, list[str]]]:
    return data.get("environments", {}).get("roomContexts", {}) or {}


def current_combat_narratives(data: dict[str, Any]) -> dict[str, list[str]]:
    return data.get("combatNarratives", {}) or {}


def apply_patch(
    data: dict[str, Any],
    location_descriptions: dict[str, list[str]],
    room_contexts: dict[str, dict[str, list[str]]],
    combat_narratives: dict[str, list[str]],
) -> None:
    ensure_nested(data)
    env = data["environments"]

    for key, lines in location_descriptions.items():
        env["locationDescriptions"][key] = list(lines)

    for biome, room_types in room_contexts.items():
        biome_bucket = env["roomContexts"].setdefault(biome, {})
        for room_type, lines in room_types.items():
            biome_bucket[room_type] = list(lines)

    for bank, lines in combat_narratives.items():
        data["combatNarratives"][bank] = list(lines)


def print_section_header(title: str, count_label: str, count: int) -> None:
    print(f"\n{title} — {count_label}: {count}")


def print_key_summary(label: str, new_lines: list[str], old_lines: list[str] | None) -> None:
    old_count = len(old_lines) if old_lines is not None else 0
    sample = truncate(new_lines[0]) if new_lines else "(empty)"
    print(f"  {label}: {len(new_lines)} lines (was {old_count})  sample: {sample!r}")


def unchanged_keys(existing: dict[str, Any], synced: dict[str, Any]) -> list[str]:
    return sorted(k for k in existing.keys() if k not in synced)


def validate_counts(
    location_descriptions: dict[str, list[str]],
    room_contexts: dict[str, dict[str, list[str]]],
    combat_narratives: dict[str, list[str]],
) -> list[str]:
    issues: list[str] = []
    if len(location_descriptions) != EXPECTED_LOCATION_DESC_KEYS:
        issues.append(
            f"locationDescriptions key count = {len(location_descriptions)} "
            f"(expected {EXPECTED_LOCATION_DESC_KEYS})"
        )
    if len(room_contexts) != EXPECTED_ROOM_CONTEXT_BIOMES:
        issues.append(
            f"roomContexts biome count = {len(room_contexts)} "
            f"(expected {EXPECTED_ROOM_CONTEXT_BIOMES})"
        )
    room_type_keys = sum(len(rt) for rt in room_contexts.values())
    if room_type_keys == 0:
        issues.append("roomContexts has no room-type keys")
    if len(combat_narratives) != EXPECTED_COMBAT_NARRATIVE_BANKS:
        issues.append(
            f"combatNarratives bank count = {len(combat_narratives)} "
            f"(expected {EXPECTED_COMBAT_NARRATIVE_BANKS})"
        )
    return issues


def print_dry_run_summary(
    xlsx_path: Path,
    sheet_name: str,
    row_count: int,
    location_descriptions: dict[str, list[str]],
    room_contexts: dict[str, dict[str, list[str]]],
    combat_narratives: dict[str, list[str]],
    data: dict[str, Any],
    parse_errors: list[str],
    validation_issues: list[str],
) -> None:
    print("=== DRY RUN (no files modified) ===")
    print(f"Source: {xlsx_path}  sheet={sheet_name}  rows={row_count}")

    old_loc = current_location_desc(data)
    print_section_header("locationDescriptions", "keys", len(location_descriptions))
    for key in sorted(location_descriptions):
        print_key_summary(key, location_descriptions[key], old_loc.get(key))

    old_room = current_room_contexts(data)
    room_type_count = sum(len(rt) for rt in room_contexts.values())
    print_section_header(
        "roomContexts",
        f"biomes {len(room_contexts)}, room-type keys",
        room_type_count,
    )
    for biome in sorted(room_contexts):
        for room_type in sorted(room_contexts[biome]):
            label = f"{biome}/{room_type}"
            new_lines = room_contexts[biome][room_type]
            old_lines = old_room.get(biome, {}).get(room_type)
            print_key_summary(label, new_lines, old_lines)

    old_combat = current_combat_narratives(data)
    print_section_header("combatNarratives", "banks", len(combat_narratives))
    for bank in sorted(combat_narratives):
        print_key_summary(bank, combat_narratives[bank], old_combat.get(bank))

    print("\nValidation:")
    if validation_issues:
        for issue in validation_issues:
            print(f"  ! {issue}")
    else:
        print(f"  ✓ locationDescriptions key count = {len(location_descriptions)}")
        print(f"  ✓ roomContexts biome count = {len(room_contexts)}")
        print(f"  ✓ combatNarratives bank count = {len(combat_narratives)}")

    if parse_errors:
        print("\nParse warnings:")
        for err in parse_errors:
            print(f"  ! {err}")

    unchanged_loc = unchanged_keys(old_loc, location_descriptions)
    unchanged_combat = unchanged_keys(old_combat, combat_narratives)
    unchanged_room_biomes = unchanged_keys(old_room, room_contexts)

    if unchanged_loc or unchanged_room_biomes or unchanged_combat:
        print("\nUnchanged JSON keys (not in sheet):")
        if unchanged_loc:
            print(f"  locationDescriptions: {', '.join(unchanged_loc)}")
        if unchanged_room_biomes:
            print(f"  roomContexts biomes: {', '.join(unchanged_room_biomes)}")
        if unchanged_combat:
            print(f"  combatNarratives: {', '.join(unchanged_combat)}")

    print("\nRun with --write to apply.")


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Sync FlavorText.json from the flavor sheet in DEMON_FIGHTER__DATA.xlsx"
    )
    parser.add_argument(
        "--xlsx",
        required=True,
        type=Path,
        help="Path to DEMON_FIGHTER__DATA.xlsx",
    )
    parser.add_argument(
        "--json",
        type=Path,
        default=DEFAULT_JSON,
        help=f"Target FlavorText.json (default: {DEFAULT_JSON.relative_to(ROOT)})",
    )
    parser.add_argument(
        "--sheet",
        default="flavor",
        help="Worksheet tab name (default: flavor)",
    )
    parser.add_argument(
        "--write",
        action="store_true",
        help="Apply changes (default is dry-run only)",
    )
    args = parser.parse_args()

    xlsx_path = args.xlsx.expanduser().resolve()
    json_path = args.json.expanduser().resolve()

    if not xlsx_path.is_file():
        print(f"Error: workbook not found: {xlsx_path}", file=sys.stderr)
        return 1
    if not json_path.is_file():
        print(f"Error: JSON not found: {json_path}", file=sys.stderr)
        return 1

    workbook = load_workbook(xlsx_path, read_only=True, data_only=True)
    try:
        sheet = find_sheet(workbook, args.sheet)
        location_descriptions, room_contexts, combat_narratives, parse_errors, row_count = (
            parse_flavor_sheet(sheet)
        )
    finally:
        workbook.close()

    data = load_json(json_path)
    validation_issues = validate_counts(location_descriptions, room_contexts, combat_narratives)

    if parse_errors:
        print("Parse issues:", file=sys.stderr)
        for err in parse_errors:
            print(f"  {err}", file=sys.stderr)

    if args.write:
        if parse_errors:
            print("Error: fix parse issues before --write", file=sys.stderr)
            return 1
        if validation_issues:
            print("Error: validation failed:", file=sys.stderr)
            for issue in validation_issues:
                print(f"  {issue}", file=sys.stderr)
            return 1

        backup_path = json_path.with_suffix(json_path.suffix + ".bak")
        shutil.copy2(json_path, backup_path)
        apply_patch(data, location_descriptions, room_contexts, combat_narratives)
        with json_path.open("w", encoding="utf-8") as f:
            json.dump(data, f, indent=2, ensure_ascii=False)
            f.write("\n")
        print(f"Wrote {json_path}")
        print(f"Backup: {backup_path}")
        return 0

    print_dry_run_summary(
        xlsx_path,
        args.sheet,
        row_count,
        location_descriptions,
        room_contexts,
        combat_narratives,
        data,
        parse_errors,
        validation_issues,
    )
    return 1 if validation_issues else 0


if __name__ == "__main__":
    sys.exit(main())
