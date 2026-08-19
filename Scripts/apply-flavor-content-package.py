#!/usr/bin/env python3
"""Apply the Demon Fighter flavor content package to the flavor sheet in the Excel workbook."""

from __future__ import annotations

import argparse
import copy
import sys
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:
    print("Error: pip install openpyxl", file=sys.stderr)
    sys.exit(1)

# --- Content package (section/bank/key -> list of text lines) ---

LOCATION_DESCRIPTIONS: dict[str, list[str]] = {
    "Castle": [
        "Cold stone corridors run longer than any single builder could have planned for.",
        "Banners hang faded past recognition, their sigils lost to whoever last cared.",
        "The keep was built in pieces, each generation patching over the last one's mistakes.",
        "Draft moves through the halls from windows no mason ever meant to leave open.",
        "Old iron gates hang half-rusted, propped rather than repaired.",
        "Torch soot climbs the walls in layers, decades stacked on decades.",
        "The floor here has been walked smooth by feet no one remembers.",
        "Somewhere above, a bell that hasn't rung in years still hangs ready.",
    ],
    "Cavern": [
        "Water drips somewhere unseen, keeping a slow count no one's listening to.",
        "The rock here has been carved by more than water, in places too regular to be natural.",
        "Cold air moves through passages that go deeper than any torch can prove.",
        "Stone columns rise where the ceiling met the floor centuries before anyone walked here.",
        "Something metallic glints where the rock has worn thin, buried deeper than curiosity reaches.",
        "The dark here has weight, thick enough to feel on bare skin.",
        "Echoes travel farther than they should, bouncing off walls that don't sound like stone.",
        "A vein of pale ore runs through the rock, warm to the touch for no reason anyone's found.",
    ],
    "Crypt": [
        "Names are carved into the walls faster than anyone could read them all.",
        "The air holds still here, undisturbed by whatever wind moves through the rest of the world.",
        "Dust settles evenly over everything, indifferent to who it was.",
        "Old bones rest in neat rows, kept that way by hands long since joined them.",
        "A cold settles into the stone that no torch quite manages to push back.",
        "Faded offerings line the lower shelves, left by people whose names went with them.",
        "Something in the deepest chamber has been sealed shut longer than the crypt's been a crypt.",
        "The silence here isn't empty — it's tended.",
    ],
    "Crystal": [
        "Light moves through the walls here in slow, deliberate pulses, keeping a rhythm of its own.",
        "The formations grow in patterns too even to be chance, layered like something built rather than grown.",
        "A faint hum rises from deep in the stone whenever you stand still long enough to notice.",
        "Every surface throws the torchlight back distorted, bent at angles that don't match the room.",
        "Warmth radiates from the crystal in places, cold in others, with no pattern anyone's mapped.",
        "Shards litter the floor, sharp enough to cut, though nothing here looks recently broken.",
        "The formations seem to lean toward the center of the room, all of them, the same way.",
        "Something beneath the crystal keeps it fed, though no one who's asked has found what.",
    ],
    "Desert": [
        "Heat rises off the sand in waves that make the horizon lie about its distance.",
        "Wind has scoured this ground smooth, uncovering things that were better left buried.",
        "Old metal juts from a dune here and there, worn thin and half-swallowed by sand.",
        "The silence out here carries farther than it should, broken only by the shift of sand settling.",
        "Glass has formed in patches on the ground, fused by heat no cookfire could produce.",
        "Bones bleach quickly under this sun, human and otherwise, indistinguishable after long enough.",
        "Something buried hums faintly beneath the dunes when the wind dies down completely.",
        "The sand here holds a shape sometimes, like something vast once lay down and never got up.",
    ],
    "Forest": [
        "Trees grow thick and old here, roots knotted around shapes too regular to be stone.",
        "Canopy hangs low enough to swallow the light whole, and something beneath it hums faint and constant.",
        "Moss has grown over everything, softening every edge until nothing keeps its original shape.",
        "The undergrowth here grows in strange, even rows, as if something once planted it on purpose.",
        "Sap runs pale along the bark, and the trees stand closer together than trees usually choose to.",
        "A clearing opens where nothing grows at all, the ground scorched flat and never healed.",
        "Water moves through the roots in channels too straight to be a stream's own doing.",
        "The deeper woods keep a silence that feels tended rather than empty.",
    ],
    "Generic": [
        "Cut stone lines every wall here, plain and unremarkable except for its age.",
        "Dust gathers in corners no one's swept in longer than anyone living remembers.",
        "The air is still, neither warm nor cold, holding whatever mood the room was built to keep.",
        "Torch brackets line the walls at even intervals, though half stand empty now.",
        "Old scorch marks mar one wall, the story behind them long since lost.",
        "The floor bears the wear of countless footsteps, none of them remembered individually.",
        "A draft moves through from somewhere unseen, carrying no particular smell.",
        "The room holds its shape well for its age, worn but not yet given up.",
    ],
    "Graveyard": [
        "Headstones lean at angles the ground itself seems to have decided on.",
        "Grass grows uneven here, thicker over some plots than others, for reasons no one asks about.",
        "The names on the oldest markers have worn past reading, though people still visit them.",
        "Fog settles low over the ground most mornings, slow to lift even at midday.",
        "A single tree stands bare year-round, its roots running deeper than any grave here.",
        "Fresh flowers sit beside markers decades old, left by hands that never knew the names.",
        "The ground here settles strangely underfoot, uneven in ways that make you watch your step.",
        "Crows keep this place more faithfully than any of the living do.",
    ],
    "Ice": [
        "Frost climbs every surface in patterns too fine to have grown naturally.",
        "The cold here settles into the bone before it settles into the air.",
        "Ice has sealed over something along one wall, its shape lost beneath the frozen sheet.",
        "Every breath hangs visible and slow to fade, longer than the cold alone explains.",
        "The floor holds a sheen of ice thick enough to walk on and thin enough to worry about.",
        "Wind finds its way through cracks too narrow to see, carrying a sound like distant breathing.",
        "Something pale is visible beneath the ice here, still and unmoving, waiting out the cold.",
        "The silence of this place is absolute, swallowed whole by the frost.",
    ],
    "Lava": [
        "Heat rolls off the stone in waves, thick enough to taste at the back of the throat.",
        "Cracks in the ground glow faint orange, breathing slow with whatever moves beneath.",
        "The air here shimmers, bending torchlight and everything else out of true.",
        "Old metal has run and reformed here, pooled into shapes that don't belong to any tool.",
        "Ash drifts down steady, fine as snow, settling over everything in a thin gray coat.",
        "The rock underfoot is warm even where the fire hasn't reached in years.",
        "Steam hisses up from unseen cracks, carrying a smell like scorched iron.",
        "Something in the deepest chamber keeps this place hot long after any flame should have died.",
    ],
    "Ruins": [
        "Walls stand here that clearly used to be something more, though no one agrees on what.",
        "Fallen columns lie half-buried, their carvings too worn to make sense of anymore.",
        "The stonework here doesn't match anything built since, cut too clean, joined too tight.",
        "Vines have claimed most of what's left, though a few surfaces resist them entirely.",
        "Something in the foundation still hums, faint enough to miss if you're not standing still.",
        "Rubble covers most of the floor, though wide sections remain oddly untouched.",
        "The proportions of this place feel wrong for people, built at a scale that doesn't fit.",
        "Whatever this once was, it was built to last far longer than it did.",
    ],
    "Swamp": [
        "Water sits still and dark here, thick enough to hide whatever's beneath it.",
        "The air hangs heavy with rot, sweet and close in a way that's hard to place.",
        "Roots twist up out of the mud in shapes that seem to reach rather than grow.",
        "Fog clings low over the water, never quite lifting even in daylight.",
        "Something metal glints beneath the surface where the water runs shallow.",
        "The ground here gives more than it should, swallowing footsteps whole.",
        "Insects hum constant and unseen, a sound that never quite stops.",
        "Old wood juts from the mud in patterns too straight to have fallen there naturally.",
    ],
    "Temple": [
        "Worn steps lead up to an altar no one alive remembers the purpose of.",
        "Old incense clings to the stone, faint but never fully gone.",
        "Carvings line the walls, worship for something whose name has long since dropped away.",
        "Light falls through a gap in the roof onto the same spot every single day.",
        "Offerings sit untouched at the altar's base, left by hands hoping for something unnamed.",
        "The quiet here feels deliberate, held rather than simply empty.",
        "Something at the center of the room draws every line of the architecture toward it.",
        "Whoever built this meant for it to be found, eventually, by someone who'd understand it.",
    ],
}

ROOM_CONTEXTS: dict[str, list[str]] = {
    "Forest/armory": [
        "Blades are racked along a wall of fused root and pale metal, none of it rusted the way it should be.",
        "Every weapon here has been kept oiled by hands that are all long gone, and it still shows.",
        "Old iron hangs from low branches, strung there by someone who wanted it kept dry and close.",
    ],
    "Forest/boss": [
        "Whatever holds this grove has been here long enough that the trees have grown around it rather than through it.",
        "The air changes at the center of this clearing, heavier, like something old is still deciding whether to notice you.",
        "Roots pull inward here, toward whatever waits at the heart of the grove, and none of them seem to mind.",
    ],
    "Forest/chamber": [
        "The walls of this hollow are packed earth and root, grown, not built, though something clearly meant for it to hold a shape.",
        "Light filters down through a gap in the canopy that always seems to land in the same place.",
        "This room breathes with the forest around it, its air thick and unmoving.",
    ],
    "Forest/kitchen": [
        "Ash from old cookfires is packed into the ground here, layered deep enough to mark generations of use.",
        "Dried herbs hang from a low beam of bent wood, forgotten by whoever hung them last.",
        "A blackened hearth sits half-swallowed by roots, still warm though no one has fed it in years.",
    ],
    "Forest/library": [
        "Bark has been peeled and etched in careful rows, a record kept by hands that trusted wood over parchment.",
        "Stacked stones mark where something was once written, worn past reading by rain and root.",
        "Whatever knowledge was kept here has gone quiet, the way old things do when no one asks after them anymore.",
    ],
    "Forest/sanctum": [
        "Something in this hollow makes every sound feel deliberate, even the ones you didn't mean to make.",
        "The trees here stand in a ring too even to be chance, and the ground within it never seems to flood.",
        "Whoever comes here comes quietly, though none of them could say exactly why.",
    ],
    "Forest/shrine": [
        "Small offerings are tucked into the roots — coin, cloth, a single carved stone worn smooth from handling.",
        "A ring of stacked rocks marks the base of the oldest tree in the grove, though no one living remembers why.",
        "Faded ribbons hang from low branches, tied by hands hoping for something they couldn't name.",
    ],
    "Forest/treasure": [
        "Something catches the light beneath the roots, buried just deep enough to have been forgotten on purpose.",
        "Coin and old metal are scattered here like they were dropped in a hurry and never retrieved.",
        "A hollow in the oldest trunk holds more than it should, packed tight and gone dry with age.",
    ],
    "Crypt/armory": [
        "Boneyard-marked blades hang here, each one earned by outliving the last hand that held it.",
        "Weapons rest beside the dead they were buried with, though more than a few have gone missing over the years.",
        "A rack of dulled steel lines the wall, kept less for use than for remembering who carried it.",
    ],
    "Crypt/boss": [
        "Whatever commands this crypt has been mistaken for one of the dead often enough that no one corrects it anymore.",
        "The deepest chamber holds something that predates every name carved into the walls above it.",
        "Even the Boneyard-trained speak carefully here, the way you'd speak around something that might still be listening.",
    ],
    "Crypt/chamber": [
        "Burial niches line every wall, stacked deeper than the crypt was ever meant to hold.",
        "The stone here has absorbed centuries of grief the way other stone absorbs rain.",
        "Some niches are marked, most aren't, and no one's kept count of which is which.",
    ],
    "Crypt/kitchen": [
        "A cold hearth sits untouched, kept only because someone once thought the dead should smell bread again.",
        "Rations meant for mourners gone stale generations ago still sit sealed on a low shelf.",
        "Whoever tended this room last left in a hurry, and no one's finished the job since.",
    ],
    "Crypt/library": [
        "Ledgers here record names faster than anyone reads them, generation stacked on generation.",
        "A scribe's careful hand gave way to a hurried one partway through these shelves, and never recovered.",
        "Old star charts hang alongside the burial records, kept by someone who thought the two belonged together.",
    ],
    "Crypt/sanctum": [
        "Whatever's kept here has outlasted every faith that used to visit it.",
        "The Boneyard order still sends someone to sit vigil, though none of them explain what for.",
        "A single unlit lamp hangs at the center, kept ready by hands that never light it.",
    ],
    "Crypt/shrine": [
        "Small tokens are pressed into the wall's cracks — coin, cloth, a lock of hair gone brittle with age.",
        "This shrine outlasted the accord that built it, tended now by habit more than belief.",
        "Someone still leaves flowers here, though the name on the marker has worn past reading.",
    ],
    "Crypt/treasure": [
        "Grave goods glitter in the low light, left with the dead by families who could spare less than they gave.",
        "Some of what's buried here is older than the crypt itself, its original owner long past guessing.",
        "A single sealed urn sits apart from the rest, untouched by even the most desperate looters.",
    ],
    "Crystal/armory": [
        "Blades here hold an edge that never seems to dull, though no one's found the reason why.",
        "Racks of gear hum faintly in time with the walls, restless even when nothing's near.",
        "Weapons stored too long in this room start to carry a faint warmth that outlasts any forge.",
    ],
    "Crystal/boss": [
        "Whatever holds this chamber has been growing into the crystal for longer than anyone's willing to guess at.",
        "The formations pulse faster the closer you get to the center, keeping time with something unseen.",
        "Every surface here seems angled toward whatever waits at the heart of the room.",
    ],
    "Crystal/chamber": [
        "The walls have grown crystal over crystal here, burying whatever shape the room used to hold.",
        "Light bends wrong in this room, throwing shadows that don't match what's casting them.",
        "Something beneath the floor keeps this chamber warmer than the rest of the formation.",
    ],
    "Crystal/kitchen": [
        "Whoever cooked here last gave up trying to keep the crystal from creeping over the hearth.",
        "A blackened pot sits fused half into the wall, more crystal now than iron.",
        "The heat in this room never quite fades, fed by something the fire alone can't account for.",
    ],
    "Crystal/library": [
        "Records here are etched into crystal slabs, the writing shifting faintly under torchlight like it's still being written.",
        "A cracked slab lies apart from the rest, its surface gone dark where the others still glow faint.",
        "Whoever kept this archive stopped adding to it generations before anyone alive was born.",
    ],
    "Crystal/sanctum": [
        "The hum here settles into your chest before you notice you're listening for it.",
        "Every crystal in this room points the same direction, patient as a compass that's stopped caring about north.",
        "Something in the center pulses slow and steady, keeping a rhythm no one's ever explained.",
    ],
    "Crystal/shrine": [
        "Small shards are set into the wall in a pattern too deliberate to be decoration.",
        "Offerings here are laid directly against the crystal, as if it might notice the difference.",
        "The light shifts faintly whenever someone kneels at this shrine, though no one claims to have caused it.",
    ],
    "Crystal/treasure": [
        "Gems catch the light here in colors that don't quite match anything found in the ground outside.",
        "A single formation stands apart from the rest, hollow at its center where something's clearly been removed.",
        "Whatever's buried beneath this crystal has been feeding it for longer than anyone's mapped.",
    ],
    "Lava/armory": [
        "Blades here have been reforged more times than anyone can count, the heat never quite letting them rest.",
        "Old metal runs down one wall in a solid sheet, weapons fused into it that no one's tried freeing.",
        "Whatever's stored here survives the heat better than it should, kept by hands that knew something about fire the rest of us don't.",
    ],
    "Lava/boss": [
        "The heat at the center of this chamber isn't natural, and neither is whatever's decided to live in it.",
        "Cracks in the floor glow brighter the deeper into the room you go, like something below is watching you approach.",
        "Whatever holds this chamber has outlasted the fire around it long enough to stop minding the heat.",
    ],
    "Lava/chamber": [
        "The walls here have run and reformed more than once, leaving the room a different shape than it started.",
        "Ash coats every surface, thick enough that footprints last for exactly as long as the next draft.",
        "Something in this room remembers being whole, though nothing left in it looks like it once was.",
    ],
    "Lava/kitchen": [
        "The hearth here never needed lighting — the floor's been hot longer than anyone's cooked on it.",
        "Old pots have melted half into slag, fused to a stone that was never meant to hold that much heat.",
        "Whoever ate here last didn't stay long enough to finish, by the look of what's left.",
    ],
    "Lava/library": [
        "What records survived here are scorched past reading, their edges curled black and brittle.",
        "A single shelf stands apart from the fire's reach, its contents untouched though no one can say why.",
        "Whatever knowledge burned here went with whoever kept it, all at once, a long time ago.",
    ],
    "Lava/sanctum": [
        "The heat here feels tended rather than accidental, like something wants it kept this way.",
        "Every crack in the floor glows the same shade, steady as a held breath.",
        "Something in the deepest part of this room has been burning since before anyone thought to ask why.",
    ],
    "Lava/shrine": [
        "Ash has been swept into careful piles here, offerings to something no one names out loud.",
        "The heat rising from this shrine feels less like fire and more like something remembering how to be warm.",
        "Old scorch marks form a shape on the wall that might be intentional, might just be time.",
    ],
    "Lava/treasure": [
        "Metal pools glitter here, cooled into shapes too deliberate to be accident.",
        "Something buried beneath the ash has kept this corner cooler than the rest of the room.",
        "What's left here survived the fire that took everything around it, and no one's figured out why.",
    ],
    "Temple/armory": [
        "Old ceremonial blades hang beside working ones, and only the dust tells you which is which anymore.",
        "Weapons here were blessed once, by someone, for some purpose no longer recorded.",
        "A single blade rests apart from the rack, kept close by whoever still tends this place.",
    ],
    "Temple/boss": [
        "Whatever holds this temple has been mistaken for its god more than once, and never bothered correcting anyone.",
        "The air changes at the center of the sanctuary, like something is still deciding whether you're welcome.",
        "Every line of this room's architecture bends toward whatever waits at its heart.",
    ],
    "Temple/chamber": [
        "Faded murals cover every wall, their meaning lost though their colors somehow haven't.",
        "The floor here is worn smoothest at its center, generations of the same path walked by the same devotion.",
        "Light falls through a gap above onto the same spot every day, like the room was built around it.",
    ],
    "Temple/kitchen": [
        "Meals prepared here were meant for pilgrims who mostly stopped coming generations ago.",
        "A blackened hearth still holds the shape of the last fire lit in it, untouched since.",
        "Dried herbs hang from the rafters, kept fresh by habit long after anyone needed them.",
    ],
    "Temple/library": [
        "Old texts line these shelves, their language older than anyone currently able to read it.",
        "A single ledger sits apart from the rest, its pages filled with careful notes on the sky.",
        "Whoever kept this archive believed the stars mattered as much as the scripture — the shelves don't separate the two.",
    ],
    "Temple/sanctum": [
        "The quiet here feels older than the temple built around it.",
        "Every surface points toward the same spot, patient as something that's been waiting a very long time.",
        "Whoever built this meant it to be found, eventually, by someone who'd understand what it was for.",
    ],
    "Temple/shrine": [
        "Offerings here span generations, some fresh, most gone to dust, all left for the same unnamed reason.",
        "A single candle burns here always, though no one admits to relighting it.",
        "The stone beneath this shrine has been worn smooth by knees, not feet.",
    ],
    "Temple/treasure": [
        "What's kept here was given, not taken — offerings collected faithfully long after anyone remembers why.",
        "A locked case sits at the center, its key lost generations before anyone thought to ask for it back.",
        "Whatever this temple was built to protect, it's still here, still waiting to be understood.",
    ],
}

COMBAT_NARRATIVES: dict[str, list[str]] = {
    "firstBlood": [
        "First blood's drawn, and the fight stops pretending to be anything else.",
        "Someone's bleeding now. That tends to settle who's serious.",
        "The first cut lands, and everyone quiets down to watch what happens next.",
        "Blood hits the ground, and the fight finds its true pace.",
        "It's no longer a standoff once the blood shows.",
    ],
    "healthRecovery": [
        "{name} steadies, breath returning, the worst of it passing for now.",
        "Something in {name} settles — not healed, just holding again.",
        "{name} shakes off the worst of it and finds their feet.",
        "Color returns to {name}'s face, and the fight isn't over yet.",
        "{name} draws a ragged breath and keeps standing.",
    ],
    "below50Percent": [
        "{name} is bleeding freely now, and still hasn't stopped moving.",
        "{name}'s breath comes ragged, but they haven't dropped their guard.",
        "Half of what {name} started with is gone, and they keep coming anyway.",
        "{name} favors one side now, the damage plain to see.",
        "{name} is hurt bad enough that it shows in every step.",
        "{name} keeps fighting through it, though the cost is written all over them.",
        "{name}'s hands shake, but they don't lower.",
        "{name} is running on less than they'd like to admit.",
    ],
    "below10Percent": [
        "{name} is barely upright, held together by stubbornness alone.",
        "One more solid hit and {name} won't be getting up.",
        "{name}'s vision swims, but they refuse to fall yet.",
        "{name} is running on fumes, and everyone in the room knows it.",
        "{name} stands only because sitting down feels like losing.",
        "Whatever's keeping {name} on their feet, it isn't much longer.",
        "{name} is a breath away from the ground.",
        "{name} won't last another exchange like the last one.",
    ],
    "criticalHit": [
        "{name} finds the gap and doesn't waste it.",
        "That one lands clean — {name} feels it connect all the way through.",
        "{name}'s strike goes exactly where it was meant to.",
        "Something gives way under {name}'s hit, and it isn't coming back from that.",
        "{name} catches the opening before it closes and makes it count.",
        "That blow from {name} lands with a weight the fight didn't have a second ago.",
        "{name} doesn't get many like that — this one was clean.",
        "The hit from {name} lands hard enough that the room goes quiet for a beat.",
    ],
    "criticalMiss": [
        "{name}'s swing goes wide, and the follow-through costs them.",
        "{name} overreaches, and the miss leaves them open.",
        "That one never had a chance — {name}'s strike finds nothing but air.",
        "{name} stumbles through the attack, off-balance and exposed.",
        "{name}'s timing slips at the worst possible moment.",
        "The strike from {name} goes nowhere useful, and they know it immediately.",
        "{name} commits to a swing that was never going to land.",
        "{name}'s attack goes wrong in a way that's going to matter in a second.",
    ],
    "environmentalAction": [
        "The room itself has an opinion, apparently: {effect}",
        "Something in the space reacts before anyone decides it should: {effect}",
        "Whatever this place is, it isn't neutral: {effect}",
        "The room joins in, uninvited: {effect}",
        "Something old in these walls stirs: {effect}",
    ],
    "healthLeadChange": [
        "The fight tips — {name} has the room's attention now.",
        "Something shifts, and {name} is the one pressing forward.",
        "The exchange turns, and {name} is no longer on the back foot.",
        "{name} finds the opening the fight's been missing.",
        "The balance moves — {name}'s the one closing distance now.",
    ],
    "intenseBattle": [
        "Neither {player} nor {enemy} is giving an inch, and both are bleeding for it.",
        "{player} and {enemy} have stopped holding anything back.",
        "This has gone past caution — {player} and {enemy} are trading everything they've got.",
        "Both {player} and {enemy} are past the point of pulling a hit.",
        "The fight between {player} and {enemy} has stopped being about winning and started being about outlasting.",
        "{player} and {enemy} are matched close enough that either outcome feels possible.",
        "There's no more room to hold back between {player} and {enemy}.",
        "{player} and {enemy} are trading blows fast enough that neither has time to think.",
    ],
    "enemyDefeated": [
        "{name} goes down, and {player} doesn't feel like celebrating it.",
        "{name} stops moving, and the quiet after is louder than the fight was.",
        "Whatever kept {name} standing runs out, and {player} is the one left breathing.",
        "{name} falls, and {player} takes a moment before moving again.",
        "It's over — {name}'s down, and {player} doesn't waste the win looking back.",
        "{name} doesn't get back up this time.",
        "{player} stands over what's left of {name}, and there's nothing satisfying about it.",
        "{name} is finished, and {player} already knows there'll be another one somewhere.",
    ],
    "playerDefeated": [
        "You go down, and {enemy} doesn't slow to check the work.",
        "The last thing you register is {enemy} still standing.",
        "You hit the ground and {enemy}'s the only thing left moving.",
        "It ends fast — {enemy} wins, and you don't get a say in how.",
        "You're out before you register the last hit landed.",
        "{enemy} finishes it, and you're not in any shape to argue.",
        "The fight's over, and you're the one on the ground this time.",
        "You go still, and {enemy} doesn't stick around to gloat.",
    ],
    "playerTaunt": [
        "\"You're slower than you think, {enemy}.\" {name} circles, watching.",
        "\"This isn't going the way you planned, is it, {enemy}?\" {name} says.",
        "\"You picked the wrong day for this, {enemy}.\" {name} keeps their guard up.",
        "\"I've had worse mornings than you, {enemy}.\" {name} mutters.",
        "\"You're going to want to reconsider this, {enemy}.\" {name} warns.",
        "\"Last chance to walk off, {enemy}.\" {name} doesn't lower their weapon.",
        "\"You're not the first to try this, {enemy}. Won't be the last either.\" {name} says flatly.",
        "\"I don't enjoy this part, {enemy}, but I'll finish it.\" {name} says.",
        "\"You're wasting both our time, {enemy}.\" {name} sighs.",
        "\"This doesn't have to go the way you think it does, {enemy}.\" {name} says, watching for an opening.",
    ],
    "enemyTaunt": [
        "\"You won't walk out of here, {player}.\" {name} says, low.",
        "\"I've buried better than you, {player}.\" {name} growls.",
        "\"You're already tired, {player}. I can see it.\" {name} watches close.",
        "\"This ends however I decide it ends, {player}.\" {name} says.",
        "\"You shouldn't have come here, {player}.\" {name} warns.",
        "\"I don't need long for this, {player}.\" {name} says, unbothered.",
        "\"You're not leaving this the way you came, {player}.\" {name} says.",
        "\"I've heard better threats than yours, {player}.\" {name} says.",
        "\"This is going to hurt more than you're ready for, {player}.\" {name} says.",
        "\"You made a mistake coming here, {player}.\" {name} says quietly.",
    ],
    "playerTaunt_forest": [
        "\"You picked the wrong grove to die in, {enemy},\" {name} mutters, backing toward the trees.",
        "\"{enemy}, this wood's older than both of us — it won't remember you long.\" {name} grins.",
        "\"Every root here's heard better threats than yours, {enemy}.\" {name} plants their feet.",
        "\"{enemy}, the trees don't take sides. Lucky for me, I don't need them to.\" {name} says.",
    ],
    "enemyTaunt_forest": [
        "\"You won't leave this grove the way you came in, {player},\" {name} growls from the shadow of the canopy.",
        "\"{player}, the roots will hold what's left of you.\" {name} circles slow.",
        "\"This wood has swallowed better than you, {player}.\" {name} snarls.",
        "\"{player}, no one hears you scream this deep in.\" {name} says, low.",
    ],
    "playerTaunt_crypt": [
        "\"The dead down here learned to stay quiet a long time ago, {enemy}. You'll learn faster.\"",
        "\"{enemy}, every door in this crypt was sealed for a reason. I'm about to be yours.\"",
        "\"Dust doesn't lie about how long something's been still.\" {name} watches {enemy}. \"You're next.\"",
        "\"{enemy}, they say the Eye doesn't close down here. Mine won't either.\"",
    ],
    "enemyTaunt_crypt": [
        "\"You'll lie still soon enough, {player}.\" {name}'s voice comes off the stone flat and close.",
        "\"{player}, this crypt's swallowed better company than you and never once complained.\"",
        "\"Every name carved into this wall outlived the hand that carved it. Yours won't.\"",
        "\"{player}, the quiet down here isn't empty. It's patient.\"",
    ],
    "playerTaunt_crystal": [
        "\"{enemy}, this room shows me a dozen versions of you. None of them win.\"",
        "\"The light lies in here, {enemy}. My blade doesn't.\"",
        "\"{enemy}, count the reflections if it helps. There's still only one of me that matters.\"",
        "\"Everything in this room is sharp if you touch it wrong. So am I.\"",
    ],
    "enemyTaunt_crystal": [
        "\"You won't see which one of me is real, {player} — not until it's too late.\"",
        "\"{player}, the light bends in here. So will you.\"",
        "\"This room's caught better reflections than yours, {player}, and kept every one.\"",
        "\"{player}, sharp edges don't ask permission. Neither do I.\"",
    ],
    "playerTaunt_lava": [
        "\"{enemy}, the ground's not the only thing about to give out from under you.\"",
        "\"Ash settles on everything down here eventually. Might as well be you first.\"",
        "\"{enemy}, mind your footing. I won't offer the warning twice.\"",
        "\"The heat doesn't care who started the fight, {enemy}. Neither do I, once it's lit.\"",
    ],
    "enemyTaunt_lava": [
        "\"You won't outlast the heat, {player}.\" Embers drift between {name} and the words.",
        "\"{player}, this ground's swallowed better than you and asked for seconds.\"",
        "\"The fire doesn't forgive slow feet, {player}, and neither will I.\"",
        "\"{player}, everything down here burns clean. You'll be no different.\"",
    ],
    "playerTaunt_temple": [
        "\"{enemy}, whatever they prayed to here stopped listening long before either of us showed up.\"",
        "\"This floor's been worn smooth by knees, not feet, {enemy}. Yours won't leave a mark.\"",
        "\"{enemy}, I don't need a blessing to finish this. Just room to swing.\"",
        "\"Every offering left in this place went unanswered, {enemy}. Don't join the pile.\"",
    ],
    "enemyTaunt_temple": [
        "\"You'll kneel here whether you meant to or not, {player}.\" {name}'s voice fills the hollow stone.",
        "\"{player}, this temple's outlasted better prayers than yours. It'll outlast you too.\"",
        "\"Whatever's owed in this place, {player}, you're about to pay part of it.\"",
        "\"{player}, nobody's listening in here. Except me.\"",
    ],
    "playerTaunt_library": [
        "\"{enemy}, half of what's written on these shelves nobody living can read. Doesn't make it less true.\"",
        "\"Every book in here outlived somebody who thought it mattered, {enemy}. So will this fight.\"",
        "\"{enemy}, I don't need to read the ending to know how this goes for you.\"",
        "\"Some things get lost between these shelves and nobody ever comes looking, {enemy}. Keep that in mind.\"",
    ],
    "enemyTaunt_library": [
        "\"You won't finish reading this one, {player}.\" Dust sifts from the shelf nearest {name}.",
        "\"{player}, this room's kept quieter things than you buried between its pages.\"",
        "\"Whatever you came looking for, {player}, it's not going to be worth what this costs.\"",
        "\"{player}, nobody comes looking for what gets lost in here. Not even you, soon enough.\"",
    ],
    "playerTaunt_underwater": [
        "\"{enemy}, the current doesn't take sides down here. Lucky for me, I don't need it to.\"",
        "\"I've held my breath longer than you'll last standing, {enemy}.\"",
        "\"{enemy}, everything down here learns to move with the water. You're running out of time to learn.\"",
        "\"The deep doesn't forgive slow hands, {enemy}, and neither will I.\"",
    ],
    "enemyTaunt_underwater": [
        "\"You won't surface from this one, {player}.\" The current pulls {name}'s words thin.",
        "\"{player}, the deep keeps what it takes. Nobody's ever gotten it back.\"",
        "\"This water's colder than it looks, {player}, and it's about to get worse for you.\"",
        "\"{player}, no one hears you scream down here. The current sees to that.\"",
    ],
    "firstBlood_nativeFauna": [
        "{name} draws blood and doesn't flinch from the smell of it — this is just what teeth and claws are for.",
        "The first cut lands, and {name} presses in harder, not softer.",
        "{name} tastes blood in the air and knows exactly what to do with that.",
    ],
    "firstBlood_feralStock": [
        "{name} draws first blood and hesitates half a step, like something in it expected a different ending to this.",
        "The cut lands, and for one strange beat {name} looks almost sorry before the fight takes over again.",
        "{name} presses the advantage, though something in the movement still remembers being called off once.",
    ],
    "firstBlood_technoEcho": [
        "{name} draws blood, and the wound it leaves doesn't spread the way a wound should.",
        "First blood, and {name}'s rhythm doesn't change at all — like the strike was always going to land.",
        "The cut opens clean and even, like {name} measured it first.",
    ],
    "criticalHit_nativeFauna": [
        "{name} finds the opening the way a predator finds a limp — instinct, not thought.",
        "That one lands with everything {name} has behind it.",
        "{name} doesn't waste the opening. Nothing wild ever does.",
    ],
    "criticalHit_feralStock": [
        "{name} strikes clean, disciplined in a way nothing truly wild should be.",
        "That hit from {name} lands with old, trained precision — like it was taught, once, by someone patient.",
        "{name}'s strike finds its mark with a control that doesn't belong to instinct alone.",
    ],
    "criticalHit_technoEcho": [
        "{name}'s strike lands exactly where it was always going to, like the outcome was decided before the swing.",
        "That hit from {name} connects with a precision nothing alive should have.",
        "{name} doesn't correct or adjust — the strike was already right the first time.",
    ],
    "criticalMiss_nativeFauna": [
        "{name} overreaches and stumbles, plain and animal about it.",
        "The swing from {name} goes wide, and for a second it looks almost embarrassed.",
        "{name} misses clean, off-balance, breathing hard.",
    ],
    "criticalMiss_feralStock": [
        "{name}'s swing goes wide, and it flinches at its own miss like it expected to be corrected.",
        "The attack from {name} stutters mid-motion, some old discipline breaking rhythm at the wrong moment.",
        "{name} overreaches, and for a beat it hesitates like it's waiting to be called off.",
    ],
    "criticalMiss_technoEcho": [
        "{name}'s strike goes wide, and the miss doesn't cost it anything it shows.",
        "The swing from {name} fails to connect, but nothing about its stance says it noticed.",
        "{name} misses, and resets to the exact same position, like nothing happened at all.",
    ],
    "below50Percent_nativeFauna": [
        "{name} is hurt bad and showing it, breathing hard, favoring one side.",
        "Half of what {name} started with is gone, and it's still coming.",
        "{name} bleeds the way anything alive bleeds, and keeps moving anyway.",
    ],
    "below50Percent_feralStock": [
        "{name} is hurt, and something in it keeps glancing toward an exit it won't take.",
        "Wounded, {name} still holds its ground with a discipline that outlasts the pain.",
        "{name} favors its injury but doesn't break formation — old training holding even now.",
    ],
    "below50Percent_technoEcho": [
        "{name} is damaged, but nothing about how it moves has actually slowed.",
        "Half of whatever {name} runs on is gone, and the rhythm hasn't changed once.",
        "{name}'s wounds don't seem to cost it anything they should.",
    ],
    "below10Percent_nativeFauna": [
        "{name} is barely standing, held up by nothing but stubbornness.",
        "One more solid hit and {name} won't get up — plain as that.",
        "{name}'s breath is ragged, legs shaking, and it still won't back down.",
    ],
    "below10Percent_feralStock": [
        "{name} is nearly finished, and something in its eyes looks almost relieved.",
        "Barely upright, {name} still holds its stance like someone's still watching.",
        "{name} is a breath from the ground, and won't break formation even now.",
    ],
    "below10Percent_technoEcho": [
        "{name} is nearly finished, and it hasn't once changed its rhythm to show it.",
        "Whatever's keeping {name} standing isn't stubbornness — it's just not done yet.",
        "{name} is a hit from falling, and its stance hasn't shifted an inch to compensate.",
    ],
    "enemyDefeated_nativeFauna": [
        "{name} goes down hard and stays down — a real, earned kill, nothing more or less.",
        "{name} doesn't get back up. It never had anything left to get up with.",
        "{player} stands over {name}, and there's nothing strange about what's left.",
    ],
    "enemyDefeated_feralStock": [
        "{name} goes still, and for a moment it looks less like a kill than something finally let go.",
        "{player} stands over {name}, and it's hard not to think of what it might have been, once.",
        "{name} doesn't get back up. Whatever it remembered, it doesn't anymore.",
    ],
    "enemyDefeated_technoEcho": [
        "{name} goes down, and whatever's inside it stops being warm faster than it should.",
        "{player} stands over {name}, and the stillness afterward is too even to be natural.",
        "{name} doesn't get back up, and what's left of it doesn't decay the way it should either.",
    ],
    "enemyTaunt_nativeFauna": [
        "\"{player}, you're bleeding and you don't even know it yet.\" {name} circles, low and patient.",
        "\"I don't need long for this, {player}.\" {name} says, unbothered, teeth bared.",
        "\"{player}, you smell like something already tired.\" {name} watches close.",
    ],
    "enemyTaunt_feralStock": [
        "\"{player}, I don't want to do this, but I will.\" {name} says, voice caught between snarl and hesitation.",
        "\"You remind me of someone I used to listen to, {player}. Doesn't change what happens next.\" {name} says.",
        "\"{player}, some part of me still waits for a hand that isn't coming. Yours won't be it.\" {name} says, low.",
    ],
    "enemyTaunt_technoEcho": [
        "\"This ends however it was always going to end, {player}.\" {name} says, voice too even to be angry.",
        "\"{player}, you're not the first. You won't be the last, and I won't remember either.\" {name} says.",
        "\"I don't tire, {player}. I don't know how.\" {name} says, flat.",
    ],
}


def build_update_lookup() -> dict[tuple[str, str, str], list[str]]:
    lookup: dict[tuple[str, str, str], list[str]] = {}
    for key, lines in LOCATION_DESCRIPTIONS.items():
        lookup[("environments", "locationDescriptions", key)] = lines
    for key, lines in ROOM_CONTEXTS.items():
        lookup[("environments", "roomContexts", key)] = lines
    for bank, lines in COMBAT_NARRATIVES.items():
        lookup[("combatNarratives", bank, "")] = lines
    return lookup


def cell_str(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def row_group_key(section: str, bank: str, key: str) -> tuple[str, str, str] | None:
    if section == "environments" and bank == "locationDescriptions":
        return (section, bank, key)
    if section == "environments" and bank == "roomContexts":
        return (section, bank, key)
    if section == "combatNarratives":
        return (section, bank, "")
    return None


def find_flavor_sheet(workbook, sheet_name: str):
    target = sheet_name.strip().lower()
    for name in workbook.sheetnames:
        if name.strip().lower() == target:
            return workbook[name]
    raise ValueError(f"Sheet {sheet_name!r} not found. Available: {workbook.sheetnames}")


def apply_updates(sheet, updates: dict[tuple[str, str, str], list[str]], dry_run: bool) -> dict[str, int]:
    header_row = 1
    header = [cell_str(sheet.cell(header_row, c).value).lower() for c in range(1, sheet.max_column + 1)]
    while header and not header[-1]:
        header.pop()
    col = {name: idx + 1 for idx, name in enumerate(header)}
    required = ("section", "bank", "key", "text")
    for name in required:
        if name not in col:
            raise ValueError(f"Missing column {name!r}")

    groups: dict[tuple[str, str, str], list[int]] = {}
    for row in range(header_row + 1, sheet.max_row + 1):
        section = cell_str(sheet.cell(row, col["section"]).value)
        bank = cell_str(sheet.cell(row, col["bank"]).value)
        key = cell_str(sheet.cell(row, col["key"]).value)
        text = cell_str(sheet.cell(row, col["text"]).value)
        if not section and not bank and not key and not text:
            continue
        gkey = row_group_key(section, bank, key)
        if gkey is None:
            continue
        groups.setdefault(gkey, []).append(row)

    stats = {"updated_keys": 0, "rows_changed": 0, "rows_added": 0, "rows_removed": 0}

    for gkey, new_lines in sorted(updates.items(), key=lambda x: groups.get(x[0], [0])[0], reverse=True):
        rows = groups.get(gkey)
        section, bank, key = gkey
        display = f"{section}/{bank}/{key}" if key else f"{section}/{bank}"
        new_count = len(new_lines)

        if not rows:
            if dry_run:
                print(f"  {display}: would ADD {new_count} lines (new bank)")
                if new_lines:
                    print(f"    sample: {new_lines[0][:72]}...")
                stats["updated_keys"] += 1
                stats["rows_added"] += new_count
                continue

            insert_at = sheet.max_row + 1
            for i, line in enumerate(new_lines):
                r = insert_at + i
                sheet.cell(r, col["section"]).value = section
                sheet.cell(r, col["bank"]).value = bank
                sheet.cell(r, col["key"]).value = key if key else None
                sheet.cell(r, col["text"]).value = line
            stats["rows_added"] += new_count
            stats["updated_keys"] += 1
            print(f"  added {display}: {new_count} lines")
            continue

        old_count = len(rows)

        if dry_run:
            print(f"  {display}: {new_count} lines (was {old_count})")
            if new_lines:
                print(f"    sample: {new_lines[0][:72]}...")
            stats["updated_keys"] += 1
            continue

        template_row = rows[0]
        if new_count > old_count:
            insert_at = rows[-1] + 1
            to_add = new_count - old_count
            sheet.insert_rows(insert_at, amount=to_add)
            for i in range(to_add):
                r = insert_at + i
                sheet.cell(r, col["section"]).value = sheet.cell(template_row, col["section"]).value
                sheet.cell(r, col["bank"]).value = sheet.cell(template_row, col["bank"]).value
                sheet.cell(r, col["key"]).value = sheet.cell(template_row, col["key"]).value
            rows = rows + list(range(insert_at, insert_at + to_add))
            stats["rows_added"] += to_add
        elif new_count < old_count:
            for r in sorted(rows[new_count:], reverse=True):
                sheet.delete_rows(r, 1)
            rows = rows[:new_count]
            stats["rows_removed"] += old_count - new_count

        for r, line in zip(rows, new_lines):
            sheet.cell(r, col["text"]).value = line
            stats["rows_changed"] += 1
        stats["updated_keys"] += 1
        print(f"  updated {display}: {new_count} lines (was {old_count})")

    return stats


def main() -> int:
    parser = argparse.ArgumentParser(description="Apply flavor content package to Excel flavor sheet")
    parser.add_argument("--xlsx", required=True, type=Path)
    parser.add_argument("--sheet", default="flavor")
    parser.add_argument("--dry-run", action="store_true", help="Show what would change without saving")
    args = parser.parse_args()

    xlsx = args.xlsx.expanduser().resolve()
    if not xlsx.is_file():
        print(f"Error: not found: {xlsx}", file=sys.stderr)
        return 1

    updates = build_update_lookup()
    print(f"{'DRY RUN' if args.dry_run else 'APPLYING'} — {len(updates)} keys in content package")
    print(f"Source: {xlsx}  sheet={args.sheet}\n")

    wb = load_workbook(xlsx)
    sheet = find_flavor_sheet(wb, args.sheet)
    stats = apply_updates(sheet, updates, dry_run=args.dry_run)

    if args.dry_run:
        print(f"\nWould update {stats['updated_keys']} keys. Run without --dry-run to save.")
        return 0

    wb.save(xlsx)
    print(
        f"\nSaved {xlsx}\n"
        f"  keys updated: {stats['updated_keys']}\n"
        f"  rows changed: {stats['rows_changed']}\n"
        f"  rows added: {stats['rows_added']}\n"
        f"  rows removed: {stats['rows_removed']}"
    )
    return 0


if __name__ == "__main__":
    sys.exit(main())
